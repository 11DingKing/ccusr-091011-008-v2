"""
仓库管理视图
"""
import logging
import io
from decimal import Decimal
from django.db import transaction
from django.db.models import Sum
from django.http import HttpResponse
from django.utils import timezone
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from apps.core.response import success_response, error_response
from .models import (
    Unit, Category, Variety, Goods, StockIn, StockOut, Warning, Approval,
    StocktakeInvestigation, InvestigationNote, StockAdjustment,
)
from .serializers import (
    UnitSerializer, UnitCreateSerializer,
    CategorySerializer, CategoryCreateSerializer,
    VarietySerializer, VarietyCreateSerializer,
    GoodsSerializer, StockInSerializer, StockOutSerializer,
    WarningSerializer, ApprovalSerializer,
    InvestigationSerializer, InvestigationCreateSerializer,
    InvestigationNoteSerializer, InvestigationNoteCreateSerializer,
    InvestigationCloseSerializer, InvestigationRevokeSerializer,
    InvestigationCancelSerializer,
)

logger = logging.getLogger('apps')


# ==================== 单位管理 ====================

class UnitListView(APIView):
    """单位列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Unit.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        units = queryset[start:end]
        
        serializer = UnitSerializer(units, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建单位"""
        serializer = UnitCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.create(
            name=serializer.validated_data['name'],
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='创建成功')


class UnitDetailView(APIView):
    """单位详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        serializer = UnitCreateSerializer(data=request.data, context={'instance': unit})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit.name = serializer.validated_data['name']
        unit.save()
        
        logger.info(f"User {request.user.username} updated unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        if unit.is_linked:
            return error_response(message='该单位已被关联，无法删除')
        
        name = unit.name
        unit.delete()
        
        logger.info(f"User {request.user.username} deleted unit {name}")
        
        return success_response(message='删除成功')


class UnitBatchDeleteView(APIView):
    """单位批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的单位')
        
        # 只删除未关联的单位
        units = Unit.objects.filter(pk__in=ids)
        deleted_count = 0
        for unit in units:
            if not unit.is_linked:
                unit.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} units")
        
        return success_response(message=f'成功删除 {deleted_count} 个单位')


class UnitAllView(APIView):
    """获取所有单位（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        units = Unit.objects.filter(is_active=True).order_by('name')
        serializer = UnitSerializer(units, many=True)
        return success_response(data=serializer.data)


# ==================== 品类管理 ====================

class CategoryListView(APIView):
    """品类列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Category.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        categories = queryset[start:end]
        
        serializer = CategorySerializer(categories, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品类"""
        serializer = CategoryCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category = Category.objects.create(
            name=serializer.validated_data['name'],
            unit=unit,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='创建成功')


class CategoryDetailView(APIView):
    """品类详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        serializer = CategoryCreateSerializer(data=request.data, context={'instance': category})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        category.name = serializer.validated_data['name']
        category.unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category.save()
        
        logger.info(f"User {request.user.username} updated category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        if category.is_linked:
            return error_response(message='该品类已被关联，无法删除')
        
        name = category.name
        category.delete()
        
        logger.info(f"User {request.user.username} deleted category {name}")
        
        return success_response(message='删除成功')


class CategoryBatchDeleteView(APIView):
    """品类批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品类')
        
        categories = Category.objects.filter(pk__in=ids)
        deleted_count = 0
        for category in categories:
            if not category.is_linked:
                category.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} categories")
        
        return success_response(message=f'成功删除 {deleted_count} 个品类')


class CategoryAllView(APIView):
    """获取所有品类（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        categories = Category.objects.filter(is_active=True).order_by('name')
        serializer = CategorySerializer(categories, many=True)
        return success_response(data=serializer.data)


# ==================== 品种管理 ====================

class VarietyListView(APIView):
    """品种列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Variety.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        varieties = queryset[start:end]
        
        serializer = VarietySerializer(varieties, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品种"""
        serializer = VarietyCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        category = Category.objects.get(pk=serializer.validated_data['category'])
        variety = Variety.objects.create(
            name=serializer.validated_data['name'],
            category=category,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='创建成功')


class VarietyDetailView(APIView):
    """品种详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        serializer = VarietyCreateSerializer(data=request.data, context={'instance': variety})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        variety.name = serializer.validated_data['name']
        variety.category = Category.objects.get(pk=serializer.validated_data['category'])
        variety.save()
        
        logger.info(f"User {request.user.username} updated variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        if variety.is_in_stock:
            return error_response(message='该品种已入库，无法删除')
        
        name = variety.name
        variety.delete()
        
        logger.info(f"User {request.user.username} deleted variety {name}")
        
        return success_response(message='删除成功')


class VarietyBatchDeleteView(APIView):
    """品种批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品种')
        
        varieties = Variety.objects.filter(pk__in=ids)
        deleted_count = 0
        for variety in varieties:
            if not variety.is_in_stock:
                variety.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} varieties")
        
        return success_response(message=f'成功删除 {deleted_count} 个品种')


class VarietyTemplateView(APIView):
    """品种导入模板下载"""
    permission_classes = []  # 允许匿名访问，通过token参数验证
    
    def get(self, request):
        # 从URL参数获取token进行验证
        from apps.authentication.backends import decode_token
        from apps.authentication.models import User
        
        token = request.query_params.get('token')
        if not token:
            return error_response(message='缺少认证信息', code=401)
        
        payload = decode_token(token)
        if not payload:
            return error_response(message='认证信息无效或已过期', code=401)
        
        try:
            user = User.objects.get(pk=payload['user_id'])
        except User.DoesNotExist:
            return error_response(message='用户不存在', code=401)
        
        wb = Workbook()
        
        # 第一个表格 - 导入模板
        ws1 = wb.active
        ws1.title = '品种导入'
        
        # 设置表头样式
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ['品种', '品类', '单位']
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 设置列宽
        ws1.column_dimensions['A'].width = 25
        ws1.column_dimensions['B'].width = 20
        ws1.column_dimensions['C'].width = 15
        
        # 第二个表格 - 品类参考
        ws2 = wb.create_sheet(title='品类参考')
        
        headers2 = ['品类', '单位']
        for col, header in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 填充品类数据
        categories = Category.objects.filter(is_active=True).select_related('unit')
        for row, category in enumerate(categories, 2):
            ws2.cell(row=row, column=1, value=category.name).border = thin_border
            ws2.cell(row=row, column=2, value=category.unit.name).border = thin_border
        
        ws2.column_dimensions['A'].width = 20
        ws2.column_dimensions['B'].width = 15
        
        # 返回Excel文件
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=variety_import_template.xlsx'
        
        return response


class VarietyImportView(APIView):
    """品种导入视图"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    
    def post(self, request):
        if 'file' not in request.FILES:
            return error_response(message='请上传文件')
        
        file = request.FILES['file']
        
        try:
            wb = load_workbook(file)
            ws = wb.active
        except Exception as e:
            return error_response(message='文件格式错误，请上传Excel文件')
        
        # 获取所有品类及其单位
        categories = {c.name: c for c in Category.objects.filter(is_active=True).select_related('unit')}
        
        can_import = []
        cannot_import = []
        
        for row in range(2, ws.max_row + 1):
            variety_name = ws.cell(row=row, column=1).value
            category_name = ws.cell(row=row, column=2).value
            unit_name = ws.cell(row=row, column=3).value
            
            if not variety_name:
                continue
            
            variety_name = str(variety_name).strip()
            category_name = str(category_name).strip() if category_name else ''
            unit_name = str(unit_name).strip() if unit_name else ''
            
            # 验证
            error_msg = None
            
            if not variety_name:
                error_msg = '品种名称不能为空'
            elif len(variety_name) > 20:
                error_msg = '品种名称最多20个字'
            elif not category_name:
                error_msg = '品类不能为空'
            elif category_name not in categories:
                error_msg = f'品类"{category_name}"不存在'
            elif not unit_name:
                error_msg = '单位不能为空'
            elif categories.get(category_name) and categories[category_name].unit.name != unit_name:
                error_msg = f'单位与品类不匹配，应为"{categories[category_name].unit.name}"'
            elif Variety.objects.filter(name=variety_name, category__name=category_name).exists():
                error_msg = '该品种已存在'
            
            if error_msg:
                cannot_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name,
                    'reason': error_msg
                })
            else:
                can_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name
                })
        
        # 如果是预览请求
        if request.data.get('preview') == 'true':
            return success_response(data={
                'can_import': can_import,
                'cannot_import': cannot_import,
                'can_import_count': len(can_import),
                'cannot_import_count': len(cannot_import)
            })
        
        # 执行导入
        imported_count = 0
        for item in can_import:
            category = categories[item['category']]
            Variety.objects.create(
                name=item['variety'],
                category=category,
                created_by=request.user
            )
            imported_count += 1
        
        logger.info(f"User {request.user.username} imported {imported_count} varieties")
        
        return success_response(
            data={
                'imported_count': imported_count,
                'failed_count': len(cannot_import),
                'failed_items': cannot_import
            },
            message=f'成功导入 {imported_count} 个品种'
        )


# ==================== 其他视图占位 ====================

class DashboardView(APIView):
    """仪表盘视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'message': '仪表盘功能开发中...'
        })


class GoodsListView(APIView):
    """货物列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class StockInListView(APIView):
    """入库记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class StockOutListView(APIView):
    """出库记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class WarningListView(APIView):
    """预警记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class ApprovalListView(APIView):
    """审批记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


# ==================== 盘点差异调查 ====================

def _generate_order_no():
    """生成调查单号：INV + 日期 + 4位当日序号"""
    prefix = f"INV{timezone.localdate():%Y%m%d}"
    last_no = StocktakeInvestigation.objects.filter(
        order_no__startswith=prefix
    ).order_by('-order_no').values_list('order_no', flat=True).first()
    seq = int(last_no[-4:]) + 1 if last_no else 1
    return f"{prefix}{seq:04d}"


def _get_investigation(pk):
    """按主键获取调查单，不存在时返回None"""
    try:
        return StocktakeInvestigation.objects.get(pk=pk)
    except StocktakeInvestigation.DoesNotExist:
        return None


class InvestigationListView(APIView):
    """盘点差异调查单列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = StocktakeInvestigation.objects.select_related(
            'goods', 'responsible_user', 'created_by', 'closed_by', 'revoked_by'
        ).prefetch_related('notes', 'adjustments').order_by('-created_at')

        status_param = request.query_params.get('status')
        if status_param:
            queryset = queryset.filter(status=status_param)
        goods_id = request.query_params.get('goods_id')
        if goods_id:
            queryset = queryset.filter(goods_id=goods_id)
        order_no = request.query_params.get('order_no')
        if order_no:
            queryset = queryset.filter(order_no__icontains=order_no)

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        serializer = InvestigationSerializer(queryset[start:end], many=True)

        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    def post(self, request):
        """创建调查单：记录盘点快照与责任范围"""
        serializer = InvestigationCreateSerializer(data=request.data)
        if not serializer.is_valid():
            first_error = list(serializer.errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))

        data = serializer.validated_data

        with transaction.atomic():
            goods = Goods.objects.select_for_update().get(pk=data['goods_id'])
            snapshot_quantity = goods.quantity
            counted_quantity = data['counted_quantity']

            if counted_quantity == snapshot_quantity:
                return error_response(message='账实相符，无需发起差异调查')

            if goods.stocktake_investigations.filter(status='investigating').exists():
                return error_response(message='该货物存在调查中的调查单，请先结案或作废')

            investigation = StocktakeInvestigation.objects.create(
                order_no=_generate_order_no(),
                goods=goods,
                responsible_user_id=data.get('responsible_user_id'),
                responsible_dept=data.get('responsible_dept', ''),
                snapshot_quantity=snapshot_quantity,
                counted_quantity=counted_quantity,
                created_by=request.user,
            )

        logger.info(
            f"User {request.user.username} created investigation "
            f"{investigation.order_no} for goods {goods.code}"
        )

        return success_response(
            data=InvestigationSerializer(investigation).data,
            message='创建成功'
        )


class InvestigationDetailView(APIView):
    """盘点差异调查单详情视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        investigation = _get_investigation(pk)
        if investigation is None:
            return error_response(message='调查单不存在', code=404)

        return success_response(data=InvestigationSerializer(investigation).data)


class InvestigationNoteListView(APIView):
    """调查过程记录视图（原因/证据/复核意见，仅可追加）"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        investigation = _get_investigation(pk)
        if investigation is None:
            return error_response(message='调查单不存在', code=404)

        if investigation.status != 'investigating':
            return error_response(message='当前状态不允许追加调查记录')

        serializer = InvestigationNoteCreateSerializer(data=request.data)
        if not serializer.is_valid():
            first_error = list(serializer.errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))

        note = InvestigationNote.objects.create(
            investigation=investigation,
            type=serializer.validated_data['type'],
            content=serializer.validated_data['content'],
            created_by=request.user,
        )

        logger.info(
            f"User {request.user.username} added {note.type} note "
            f"to investigation {investigation.order_no}"
        )

        return success_response(
            data=InvestigationNoteSerializer(note).data,
            message='记录成功'
        )


class InvestigationCloseView(APIView):
    """调查单结案视图：批准后以独立调整分录修正数量"""
    permission_classes = [IsAuthenticated]

    REQUIRED_NOTE_TYPES = [('reason', '原因分析'), ('evidence', '证据材料'), ('review', '复核意见')]

    def post(self, request, pk):
        if not request.user.is_admin:
            return error_response(message='仅管理员可以批准结案', code=403)

        if _get_investigation(pk) is None:
            return error_response(message='调查单不存在', code=404)

        serializer = InvestigationCloseSerializer(data=request.data)
        if not serializer.is_valid():
            first_error = list(serializer.errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))

        with transaction.atomic():
            investigation = StocktakeInvestigation.objects.select_for_update().get(pk=pk)

            if investigation.status != 'investigating':
                return error_response(message='当前状态不允许结案')

            if investigation.created_by_id == request.user.id:
                return error_response(message='调查单创建人不能批准自己的调查单')

            existing_types = set(investigation.notes.values_list('type', flat=True))
            missing = [label for note_type, label in self.REQUIRED_NOTE_TYPES
                       if note_type not in existing_types]
            if missing:
                return error_response(message=f"结案前需补充：{'、'.join(missing)}")

            goods = Goods.objects.select_for_update().get(pk=investigation.goods_id)

            # 结案时检测基线是否已变化，并按快照重新计算差额
            current_quantity = goods.quantity
            baseline_changed = current_quantity != investigation.snapshot_quantity
            difference = investigation.counted_quantity - investigation.snapshot_quantity
            after_quantity = current_quantity + difference

            if after_quantity < 0:
                return error_response(message='调整后库存不能为负，请核实后续收发记录')

            StockAdjustment.objects.create(
                investigation=investigation,
                goods=goods,
                type='stocktake',
                quantity=difference,
                before_quantity=current_quantity,
                after_quantity=after_quantity,
                operator=request.user,
                remark=f'盘点差异调整（调查单 {investigation.order_no}）',
            )
            goods.quantity = after_quantity
            goods.save(update_fields=['quantity', 'updated_at'])

            investigation.status = 'closed'
            investigation.conclusion = serializer.validated_data['conclusion']
            investigation.close_baseline_quantity = current_quantity
            investigation.baseline_changed = baseline_changed
            investigation.closed_by = request.user
            investigation.closed_at = timezone.now()
            investigation.save()

        logger.info(
            f"User {request.user.username} closed investigation "
            f"{investigation.order_no}, adjustment {difference}"
        )

        return success_response(
            data=InvestigationSerializer(investigation).data,
            message='结案成功'
        )


class InvestigationCancelView(APIView):
    """调查单作废视图（仅调查中可作废，不影响库存）"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        if _get_investigation(pk) is None:
            return error_response(message='调查单不存在', code=404)

        serializer = InvestigationCancelSerializer(data=request.data)
        if not serializer.is_valid():
            first_error = list(serializer.errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))

        with transaction.atomic():
            investigation = StocktakeInvestigation.objects.select_for_update().get(pk=pk)

            if investigation.status != 'investigating':
                return error_response(message='仅调查中的调查单可以作废')

            investigation.status = 'cancelled'
            investigation.cancel_reason = serializer.validated_data.get('reason', '')
            investigation.save(update_fields=['status', 'cancel_reason'])

        logger.info(
            f"User {request.user.username} cancelled investigation "
            f"{investigation.order_no}"
        )

        return success_response(
            data=InvestigationSerializer(investigation).data,
            message='作废成功'
        )


class InvestigationRevokeView(APIView):
    """调查结论撤销视图：以冲回分录抵销，不覆盖旧记录"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        if not request.user.is_admin:
            return error_response(message='仅管理员可以撤销调查结论', code=403)

        if _get_investigation(pk) is None:
            return error_response(message='调查单不存在', code=404)

        serializer = InvestigationRevokeSerializer(data=request.data)
        if not serializer.is_valid():
            first_error = list(serializer.errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))

        with transaction.atomic():
            investigation = StocktakeInvestigation.objects.select_for_update().get(pk=pk)

            if investigation.status != 'closed':
                return error_response(message='仅已结案的调查单可以撤销')

            original = investigation.adjustments.filter(type='stocktake').first()
            if original is None:
                return error_response(message='结案调整分录缺失，无法撤销', code=500)

            goods = Goods.objects.select_for_update().get(pk=investigation.goods_id)

            current_quantity = goods.quantity
            reversal_quantity = -original.quantity
            after_quantity = current_quantity + reversal_quantity

            if after_quantity < 0:
                return error_response(message='冲回后库存不能为负，无法撤销')

            StockAdjustment.objects.create(
                investigation=investigation,
                goods=goods,
                type='reversal',
                quantity=reversal_quantity,
                before_quantity=current_quantity,
                after_quantity=after_quantity,
                reverses=original,
                operator=request.user,
                remark=f'撤销冲回（调查单 {investigation.order_no}）',
            )
            goods.quantity = after_quantity
            goods.save(update_fields=['quantity', 'updated_at'])

            investigation.status = 'revoked'
            investigation.revoke_reason = serializer.validated_data['reason']
            investigation.revoked_by = request.user
            investigation.revoked_at = timezone.now()
            investigation.save()

        logger.info(
            f"User {request.user.username} revoked investigation "
            f"{investigation.order_no}"
        )

        return success_response(
            data=InvestigationSerializer(investigation).data,
            message='撤销成功'
        )


class InvestigationReconciliationView(APIView):
    """差异还原视图：还原发现值、系统值、后续变动与实际调整的关系"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        investigation = _get_investigation(pk)
        if investigation is None:
            return error_response(message='调查单不存在', code=404)

        goods = investigation.goods
        snapshot_at = investigation.snapshot_at

        # 快照之后的后续收发变动
        stock_ins = StockIn.objects.filter(
            goods=goods, stock_in_time__gte=snapshot_at
        ).select_related('operator')
        stock_outs = StockOut.objects.filter(
            goods=goods, status='completed', stock_out_time__gte=snapshot_at
        ).select_related('operator')

        in_total = stock_ins.aggregate(total=Sum('quantity'))['total'] or Decimal('0')
        out_total = stock_outs.aggregate(total=Sum('quantity'))['total'] or Decimal('0')

        movements = []
        for record in stock_ins:
            movements.append({
                'type': 'in',
                'type_display': '入库',
                'quantity': record.quantity,
                'time': record.stock_in_time,
                'operator': record.operator.username if record.operator else '',
                'detail': record.batch_no or record.supplier,
            })
        for record in stock_outs:
            movements.append({
                'type': 'out',
                'type_display': '出库',
                'quantity': -record.quantity,
                'time': record.stock_out_time,
                'operator': record.operator.username if record.operator else '',
                'detail': record.receiver,
            })
        movements.sort(key=lambda item: item['time'])

        adjustments = investigation.adjustments.all()
        adjustment_total = sum(
            (item.quantity for item in adjustments), Decimal('0')
        )

        net_change = in_total - out_total
        expected_quantity = investigation.snapshot_quantity + net_change + adjustment_total
        current_quantity = goods.quantity

        return success_response(data={
            'order_no': investigation.order_no,
            'goods_id': goods.id,
            'goods_name': goods.name,
            'goods_code': goods.code,
            'status': investigation.status,
            'status_display': investigation.get_status_display(),
            'snapshot_at': snapshot_at,
            'snapshot_quantity': investigation.snapshot_quantity,
            'counted_quantity': investigation.counted_quantity,
            'difference': investigation.difference,
            'subsequent_movements': {
                'stock_in_total': in_total,
                'stock_out_total': out_total,
                'net_change': net_change,
                'records': movements,
            },
            'close_baseline_quantity': investigation.close_baseline_quantity,
            'baseline_changed': investigation.baseline_changed,
            'adjustments': [
                {
                    'id': item.id,
                    'type': item.type,
                    'type_display': item.get_type_display(),
                    'quantity': item.quantity,
                    'before_quantity': item.before_quantity,
                    'after_quantity': item.after_quantity,
                    'reverses': item.reverses_id,
                    'operator': item.operator.username if item.operator else '',
                    'created_at': item.created_at,
                }
                for item in adjustments
            ],
            'adjustment_total': adjustment_total,
            'current_quantity': current_quantity,
            'expected_quantity': expected_quantity,
            'is_balanced': expected_quantity == current_quantity,
        })
